"""Excel report laid out after the NSW embodied-carbon reporting structure.

Sheets: Summary (tCO2e by life-cycle module and per declared unit), Line items
(element, module, resource group, data-quality tier, basis), Strategies,
Factors (every factor used, with source and tier), Inputs, Warnings.

NOTE: column layout follows the structure described in the NSW Embodied Carbon
Measurement Technical Guide (Apr 2025); map to the official Appendix 9 template
once it is obtained.
"""
from __future__ import annotations

import io
import json
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .carbon import Result
from .factors import FactorLibrary
from .models import ProjectInput

HDR = PatternFill("solid", fgColor="1F3A5F")
HFONT = Font(color="FFFFFF", bold=True)
STAGE = {"business_case": "Stage 1 - Business case", "design": "Stage 2 - Planning / design / procurement",
         "construction": "Stage 3 - Construction / practical completion"}


def _header(ws, row, cols):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.fill, cell.font = HDR, HFONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, widths=None):
    for j, col in enumerate(ws.columns, 1):
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, w + 2), 70)
    if widths:
        for k, v in widths.items():
            ws.column_dimensions[k].width = v


def build_workbook(p: ProjectInput, res: Result, lib: FactorLibrary, strat: dict | None = None,
                   route: dict | None = None) -> bytes:
    wb = Workbook()
    d = res.to_dict()
    L = res.length_m

    ws = wb.active; ws.title = "Summary"
    ws["A1"] = "TunCO2 Pro - Upfront embodied carbon report (EN 15978 A1-A5)"; ws["A1"].font = Font(bold=True, size=14)
    meta = [("Project", p.name), ("Reporting stage", STAGE[p.stage]), ("Asset type", "Linear infrastructure - bored tunnel (TBM)"),
            ("Tunnel length assessed (m)", L), ("Declared unit", "tCO2e per route-km (single bore)"),
            ("Report date", date.today().isoformat()), ("Factor library", res.factor_library_version),
            ("Calculation mode", "v1 compatibility" if res.compat_v1 else "Corrected (default)"),
            ("Level of detail", f"LoD {p.lod}"),
            ("Functional unit", f"{p.functional.kind}, {p.functional.count} {'track(s)' if p.functional.kind == 'rail' else 'lane(s)' if p.functional.kind == 'road' else 'bore'}")]
    for i, (k, v) in enumerate(meta, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True); ws.cell(row=i, column=2, value=v)
    r0 = 3 + len(meta) + 1
    _header(ws, r0, ["Life-cycle module", "tCO2e (total)", "tCO2e per route-km", "Share"])
    tot = d["total_tCO2e"]
    for i, (m, t) in enumerate(d["modules_tCO2e"].items(), r0 + 1):
        ws.cell(row=i, column=1, value=m)
        ws.cell(row=i, column=2, value=round(t, 1))
        ws.cell(row=i, column=3, value=round(d["modules_kgCO2e_per_m"][m], 1))
        ws.cell(row=i, column=4, value=round(t / tot, 4) if tot else 0).number_format = "0.0%"
    rt = r0 + 4
    ws.cell(row=rt, column=1, value="Total A1-A5").font = Font(bold=True)
    ws.cell(row=rt, column=2, value=round(tot, 1)).font = Font(bold=True)
    ws.cell(row=rt, column=3, value=round(d["total_kgCO2e_per_m"], 1)).font = Font(bold=True)
    rf = rt + 2
    ws.cell(row=rf, column=1, value="Functional-unit normalisation (A1-A5)").font = Font(bold=True)
    for i, (k, v) in enumerate(res.functional_units().items(), rf + 1):
        ws.cell(row=i, column=1, value=k); ws.cell(row=i, column=2, value=round(v, 2))
    _autosize(ws, {"A": 44, "B": 44})

    ws = wb.create_sheet("Line items")
    _header(ws, 1, ["Module", "Element", "Resource group", "kgCO2e per m", "tCO2e total", "Data-quality tier (NSW EF hierarchy)", "Basis"])
    for i, it in enumerate(res.items, 2):
        ws.append([it.module, it.element, it.material_group, round(it.kg_per_m, 3),
                   round(it.kg_per_m * L / 1000, 3), it.nsw_tier, it.basis])
    _autosize(ws)

    if strat:
        ws = wb.create_sheet("Strategies")
        _header(ws, 1, ["Decarbonisation lever", "Saving kgCO2e per m", "Saving tCO2e total", "Saving % of total"])
        for k, v in strat["savings_kgCO2e_per_m"].items():
            ws.append([k, round(v, 1), round(v * L / 1000, 1), round(v / d["total_kgCO2e_per_m"], 4)])
            ws.cell(row=ws.max_row, column=4).number_format = "0.0%"
        if "combined_design_levers_kgCO2e_per_m" in strat:
            c = strat["combined_design_levers_kgCO2e_per_m"]
            ws.append(["Combined design levers (sequential)", round(c, 1), round(c * L / 1000, 1), round(c / d["total_kgCO2e_per_m"], 4)])
            ws.cell(row=ws.max_row, column=4).number_format = "0.0%"
        ws.append([]); ws.append([strat["note"]])
        for c in strat["checks"]:
            ws.append([c])
        _autosize(ws)

    if route:
        ws = wb.create_sheet("Route")
        ws["A1"] = f"Alignment {route['alignment']['name']}: {route['alignment']['length_m']:,.0f} m, {route['rings']:,} rings"
        _header(ws, 3, ["Zone", "Ch from", "Ch to", "Length m", "Cover m", "p0 MPa", "t m", "f'c MPa",
                        "kgCO2e/m", "tCO2e", "FoS", "u_eq mm", "ULS", "SLS"])
        for z in route["zones"]:
            ws.append([z["zone"], z["ch_from"], z["ch_to"], z["length_m"], z["cover_m"], round(z["p0_mpa"], 3),
                       z["thickness_m"], z["strength_mpa"], round(z["kgCO2e_per_m"], 1), round(z["tCO2e"], 1),
                       round(z["fos"], 2) if z["fos"] else "inf", round(z["u_mob_mm"], 1),
                       "OK" if z["uls_ok"] else "FAIL", "OK" if z["sls_ok"] else "FAIL"])
        ws.append(["Route total", None, None, None, None, None, None, None,
                   round(route["average_kgCO2e_per_m"], 1), round(route["total_tCO2e"], 1)])
        for w in route["warnings"]:
            ws.append([w])
        _autosize(ws)

    ws = wb.create_sheet("Factors")
    _header(ws, 1, ["Key", "Category", "Name", "Value", "Unit", "Source", "NSW tier", "Notes"])
    for f in lib.all():
        ws.append([f.key, f.category, f.name, f.value, f.unit, f.source, f.nsw_tier, f.notes])
    _autosize(ws)

    ws = wb.create_sheet("Inputs")
    ws["A1"] = "Full input record (JSON) for audit and re-run"
    ws["A2"] = json.dumps(p.model_dump(), indent=1)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Warnings")
    _header(ws, 1, ["Warning / assumption"])
    for w in res.warnings or ["None"]:
        ws.append([w])
    _autosize(ws)

    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
